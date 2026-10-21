"""
Tests de cobertura profunda del motor de reportes (reports.py y report_docx.py).
Fase DoD-8 / Issue #6.
"""
from __future__ import annotations

import os
from datetime import date, datetime, time
from unittest.mock import patch

import pytest
from docx import Document
from docx.shared import RGBColor

from app.domain.report_docx import (
    _add_heading,
    _add_rule,
    _cell_text,
    _fmt_fecha,
    _rgb_hex,
    _seccion_almuerzo,
    _seccion_detalle_persona,
    _seccion_horas_contrato_docx,
    _seccion_incompletos,
    _seccion_lista_tardanza,
    _seccion_salidas_anticipadas,
    _set_cell_bg,
    _set_cell_borders,
    _setup_document,
    _tabla_datos,
    _tabla_portada,
    generar_docx,
    generar_docx_persona,
)
from app.domain.reports import (
    DEFAULT_CONFIG,
    DEFAULT_FILTROS,
    _fmt_horas,
    _retraso_de_obs,
    analizar_dia,
    analizar_por_persona,
    build_pdf,
    enriquecer_analisis_con_origen,
    generar_pdf,
    generar_pdf_persona,
)


@pytest.fixture
def mock_horarios_standard():
    """Horarios estándar de prueba lunes-viernes 08:00-17:00 con 60m almuerzo."""
    return {
        "by_id": {
            "101": {
                "nombre": "Carlos Perez",
                "lunes": "08:00", "lunes_salida": "17:00",
                "martes": "08:00", "martes_salida": "17:00",
                "miercoles": "08:00", "miercoles_salida": "17:00",
                "jueves": "08:00", "jueves_salida": "17:00",
                "viernes": "08:00", "viernes_salida": "17:00",
                "sabado": None,
                "domingo": None,
                "almuerzo_min": 60,
                "horas_semana": 40.0,
                "horas_mes": 160.0,
            },
            "102": {
                "nombre": "Maria Gomez",
                "lunes": "09:00", "lunes_salida": "18:00",
                "martes": "09:00", "martes_salida": "18:00",
                "miercoles": "09:00", "miercoles_salida": "18:00",
                "jueves": "09:00", "jueves_salida": "18:00",
                "viernes": "09:00", "viernes_salida": "18:00",
                "sabado": "09:00", "sabado_salida": "13:00",
                "domingo": None,
                "almuerzo_min": 60,
                "horas_semana": 44.0,
                "horas_mes": 176.0,
            },
        },
        "by_nombre": {
            "CARLOS PEREZ": {
                "nombre": "Carlos Perez",
                "lunes": "08:00", "lunes_salida": "17:00",
                "martes": "08:00", "martes_salida": "17:00",
                "miercoles": "08:00", "miercoles_salida": "17:00",
                "jueves": "08:00", "jueves_salida": "17:00",
                "viernes": "08:00", "viernes_salida": "17:00",
                "sabado": None,
                "domingo": None,
                "almuerzo_min": 60,
                "horas_semana": 40.0,
                "horas_mes": 160.0,
            },
        },
    }


def _make_punch(nombre, id_user, fecha_str, hora_str, tipo):
    d = date.fromisoformat(fecha_str)
    t = time.fromisoformat(hora_str)
    dt = datetime.combine(d, t)
    return {
        "nombre": nombre,
        "id_usuario": id_user,
        "fecha": d,
        "datetime": dt,
        "hora": t,
        "tipo": tipo,
    }


# ==============================================================================
# 1. ANALIZAR_DIA
# ==============================================================================
class TestAnalizarDiaCoverage:

    def test_analizar_dia_completo(self, mock_horarios_standard):
        f = "2026-07-06"  # Lunes
        feriados = {date(2026, 7, 7)}
        justificaciones = {
            ("101", f, "tardanza"): {"motivo": "Cita médica"},
            ("101", f, "almuerzo"): {"motivo": "Reunión de trabajo"},
        }

        registros = [
            # Carlos: normal con justificaciones probadas
            _make_punch("Carlos Perez", "101", f, "08:04", "Entrada"),  # tardanza leve pero justificada
            _make_punch("Carlos Perez", "101", f, "12:00", "Salida"),
            _make_punch("Carlos Perez", "101", f, "13:20", "Entrada"),  # exceso almuerzo pero justificado
            _make_punch("Carlos Perez", "101", f, "17:00", "Salida"),

            # Maria: tardanza severa y salida anticipada severa sin justificar
            _make_punch("Maria Gomez", "102", f, "09:20", "Entrada"),  # +20 min severa
            _make_punch("Maria Gomez", "102", f, "12:30", "Salida"),
            _make_punch("Maria Gomez", "102", f, "14:00", "Entrada"),  # 90m almuerzo vs 60m
            _make_punch("Maria Gomez", "102", f, "17:30", "Salida"),  # salida 30m antes (severa)

            # Persona sin horario (con y sin permiso)
            _make_punch("Pedro SinHorario", "999", f, "08:00", "Entrada"),

            # Marcación en día libre
            _make_punch("Carlos Perez", "101", "2026-07-12", "10:00", "Entrada"),  # Domingo libre
        ]

        res = analizar_dia(
            registros,
            mock_horarios_standard,
            justificaciones=justificaciones,
            feriados=feriados,
            permitir_sin_horario=True,
        )

        assert "tardanza_severa" in res
        assert any(t["nombre"] == "Maria Gomez" for t in res["tardanza_severa"])
        assert "almuerzo_largo" in res
        assert any(a["nombre"] == "Maria Gomez" for a in res["almuerzo_largo"])
        assert "salida_anticipada_severa" in res
        assert any(s["nombre"] == "Maria Gomez" for s in res["salida_anticipada_severa"])
        assert "registros_incompletos" in res
        assert any("Pedro SinHorario" in inc["nombre"] for inc in res["registros_incompletos"])

    def test_analizar_dia_tardanza_leve_y_salida_leve(self, mock_horarios_standard):
        f = "2026-07-06"
        registros = [
            _make_punch("Carlos Perez", "101", f, "08:03", "Entrada"),  # +3 min leve
            _make_punch("Carlos Perez", "101", f, "16:56", "Salida"),   # 4 min antes leve
        ]
        res = analizar_dia(registros, mock_horarios_standard)
        assert any(t["nombre"] == "Carlos Perez" for t in res["tardanza_leve"])
        assert any(s["nombre"] == "Carlos Perez" for s in res["salida_anticipada_leve"])

    def test_analizar_dia_en_feriado_ignora(self, mock_horarios_standard):
        f = "2026-07-06"
        registros = [_make_punch("Carlos Perez", "101", f, "08:30", "Entrada")]
        res = analizar_dia(registros, mock_horarios_standard, feriados={date(2026, 7, 6)})
        assert len(res["tardanza_severa"]) == 0


# ==============================================================================
# 2. ANALIZAR_POR_PERSONA
# ==============================================================================
class TestAnalizarPorPersonaCoverage:

    def test_analizar_por_persona_escenarios_completos(self, mock_horarios_standard):
        justificaciones = {
            # Tardanza autorizada hasta cierta hora
            ("101", "2026-07-07", "tardanza"): {
                "hora_permitida": "08:30",
                "motivo": "Tráfico puente",
                "recuperable": 1,
                "fecha_recuperacion": "2026-07-08",
                "hora_recuperacion": "17:00",
                "hora_recuperacion_fin": "17:30",
            },
            # Tardanza autorizada sobrepasada (severa)
            ("101", "2026-07-08", "tardanza"): {
                "hora_permitida": "08:15",
                "motivo": "Cita banco",
            },
            # Permiso aprobado con retorno a tiempo
            ("101", "2026-07-09", "permiso"): {
                "estado": "aprobada",
                "hora_desde": "10:00",
                "hora_hasta": "11:00",
                "motivo": "Trámite IESS",
            },
            # Permiso sin retorno
            ("101", "2026-07-10", "permiso"): {
                "estado": "aprobada",
                "hora_desde": "14:00",
                "hora_hasta": "17:00",
                "sin_retorno": True,
                "motivo": "Calamidad doméstica",
            },
            # Justificación de anomalía / incompleto
            ("102", "2026-07-06", "incompleto"): {
                "motivo": "Olvidó registrar salida",
            },
        }

        feriados = {date(2026, 7, 11)}
        breaks_categorizados = {
            "102": {
                date(2026, 7, 7): [
                    {"inicio": "10:00", "fin": "10:15", "tipo": "pausa"},
                    {"inicio": "15:00", "fin": "15:15", "tipo": "pausa"},
                ]
            }
        }

        registros = [
            # Carlos: Día 1 (2026-07-06) normal
            _make_punch("Carlos Perez", "101", "2026-07-06", "08:00", "Entrada"),
            _make_punch("Carlos Perez", "101", "2026-07-06", "17:00", "Salida"),

            # Carlos: Día 2 (2026-07-07) llegada 08:20 autorizada hasta 08:30 (recuperable)
            _make_punch("Carlos Perez", "101", "2026-07-07", "08:20", "Entrada"),
            _make_punch("Carlos Perez", "101", "2026-07-07", "17:00", "Salida"),

            # Carlos: Día 3 (2026-07-08) llegada 08:25 excede autorizada 08:15 (+10m severa)
            _make_punch("Carlos Perez", "101", "2026-07-08", "08:25", "Entrada"),
            _make_punch("Carlos Perez", "101", "2026-07-08", "17:00", "Salida"),

            # Carlos: Día 4 (2026-07-09) permiso 10:00 - 11:00 con retorno a tiempo
            _make_punch("Carlos Perez", "101", "2026-07-09", "08:00", "Entrada"),
            _make_punch("Carlos Perez", "101", "2026-07-09", "10:00", "Salida"),
            _make_punch("Carlos Perez", "101", "2026-07-09", "10:55", "Entrada"),
            _make_punch("Carlos Perez", "101", "2026-07-09", "17:00", "Salida"),

            # Carlos: Día 5 (2026-07-10) permiso sin retorno desde las 14:00
            _make_punch("Carlos Perez", "101", "2026-07-10", "08:00", "Entrada"),
            _make_punch("Carlos Perez", "101", "2026-07-10", "14:00", "Salida"),

            # Carlos: Día libre trabajado (Sábado 2026-07-11 feriado + extra)
            _make_punch("Carlos Perez", "101", "2026-07-12", "09:00", "Entrada"),
            _make_punch("Carlos Perez", "101", "2026-07-12", "13:00", "Salida"),

            # Maria: Día 1 incompleto justificado
            _make_punch("Maria Gomez", "102", "2026-07-06", "09:00", "Entrada"),

            # Maria: Día 2 multi-breaks categorizados (6 marcaciones)
            _make_punch("Maria Gomez", "102", "2026-07-07", "09:00", "Entrada"),
            _make_punch("Maria Gomez", "102", "2026-07-07", "10:00", "Salida"),
            _make_punch("Maria Gomez", "102", "2026-07-07", "10:15", "Entrada"),
            _make_punch("Maria Gomez", "102", "2026-07-07", "15:00", "Salida"),
            _make_punch("Maria Gomez", "102", "2026-07-07", "15:15", "Entrada"),
            _make_punch("Maria Gomez", "102", "2026-07-07", "18:00", "Salida"),

            # Maria: Día 3 multi-breaks sin categorizar (pendiente revisión)
            _make_punch("Maria Gomez", "102", "2026-07-08", "09:00", "Entrada"),
            _make_punch("Maria Gomez", "102", "2026-07-08", "11:00", "Salida"),
            _make_punch("Maria Gomez", "102", "2026-07-08", "11:30", "Entrada"),
            _make_punch("Maria Gomez", "102", "2026-07-08", "14:00", "Salida"),
            _make_punch("Maria Gomez", "102", "2026-07-08", "15:00", "Entrada"),
            _make_punch("Maria Gomez", "102", "2026-07-08", "18:00", "Salida"),

            # Maria: Día 4 salida anticipada justificada
            _make_punch("Maria Gomez", "102", "2026-07-09", "09:00", "Entrada"),
            _make_punch("Maria Gomez", "102", "2026-07-09", "16:00", "Salida"),

            # Sin horario persona
            _make_punch("Sin Horario", "888", "2026-07-06", "08:00", "Entrada"),
            _make_punch("Sin Horario", "888", "2026-07-06", "17:00", "Salida"),
        ]

        # Justificación de salida anticipada para Maria el 07-09
        justificaciones[("102", "2026-07-09", "salida_anticipada")] = {"motivo": "Cita médica"}

        analisis = analizar_por_persona(
            registros=registros,
            config={},
            horarios=mock_horarios_standard,
            fecha_inicio=date(2026, 7, 6),
            fecha_fin=date(2026, 7, 12),
            justificaciones=justificaciones,
            feriados=feriados,
            breaks_categorizados=breaks_categorizados,
            mostrar_todos=True,
            permitir_sin_horario=True,
            verificar_horas=True,
            mostrar_tiempo_extra=True,
        )

        assert "Carlos Perez" in analisis
        assert "Maria Gomez" in analisis
        assert "Sin Horario" in analisis

        carlos = analisis["Carlos Perez"]
        assert carlos["resumen"]["justificadas"] >= 1
        assert any(d.get("estado") == "extra" for d in carlos["dias"])

        maria = analisis["Maria Gomez"]
        assert any(d.get("estado") == "pendiente_revision" for d in maria["dias"])

    def test_enriquecer_analisis_con_origen(self):
        analisis = {
            "Carlos Perez": {
                "dias": [
                    {"fecha": "2026-07-06", "estado": "ok"},
                ]
            }
        }
        regs_persona = {
            "Carlos Perez": [
                {"persona_id": "p-123", "fecha": "2026-07-06"}
            ]
        }
        with patch("app.domain.horarios_resolucion.resolver_horario_vigente_para_persona") as mock_res:
            mock_res.return_value = {"origen": "personalizado", "plantilla_id": "pl-1"}
            res = enriquecer_analisis_con_origen(analisis, regs_persona)
            assert res["Carlos Perez"]["dias"][0]["horario_origen"] == "personalizado"


# ==============================================================================
# 3. RENDERING PDF (ReportLab)
# ==============================================================================
class TestGenerarPDFCoverage:

    def test_generar_pdf_general(self, tmp_path):
        out_pdf = str(tmp_path / "general.pdf")
        analisis_por_dia = {
            date(2026, 7, 6): {
                "tardanza_severa": [{"nombre": "Carlos Perez", "hora": "08:20", "retraso": 20, "programado": "08:00"}],
                "tardanza_leve": [{"nombre": "Maria Gomez", "hora": "09:04", "retraso": 4, "programado": "09:00"}],
                "almuerzo_largo": [{"nombre": "Carlos Perez", "salida": "12:00", "regreso": "13:30", "duracion": 90, "exceso": 30}],
                "salida_anticipada_severa": [{"nombre": "Maria Gomez", "hora": "16:00", "programado": "18:00", "retraso": 120}],
                "salida_anticipada_leve": [],
                "registros_incompletos": [{"nombre": "Pedro", "registros": 1, "detalle": "Entrada 08:00"}],
                "resumen": {
                    "total_personas": 3,
                    "tardanza_severa": 1,
                    "tardanza_leve": 1,
                    "almuerzo_largo": 1,
                    "incompletos": 1,
                    "salida_anticipada_severa": 1,
                    "salida_anticipada_leve": 0,
                },
            }
        }
        log_dup = [
            {"nombre": "Carlos", "fecha": date(2026, 7, 6), "hora_orig": "08:00", "hora_dup": "08:00", "tipo": "Entrada", "diferencia": 0.1}
        ]

        generar_pdf(
            ruta_salida=out_pdf,
            analisis_por_dia=analisis_por_dia,
            log_duplicados=log_dup,
            config=DEFAULT_CONFIG,
            nombre_archivo_origen="test.csv",
            filtros=DEFAULT_FILTROS,
            sin_horario=["Pedro"],
        )

        assert os.path.exists(out_pdf)
        assert os.path.getsize(out_pdf) > 0

    def test_generar_pdf_persona(self, tmp_path):
        out_pdf = str(tmp_path / "persona.pdf")
        analisis_persona = {
            "Carlos Perez": {
                "dias": [
                    {
                        "fecha": date(2026, 7, 6),
                        "llegada": "08:20",
                        "salida": "17:00",
                        "hora_programada": "08:00",
                        "almuerzo_duracion": "60m",
                        "almuerzo_salida": "12:00",
                        "almuerzo_regreso": "13:00",
                        "almuerzo_exceso": None,
                        "tiempo_dentro": "8h 40m",
                        "n_registros": 4,
                        "detalle_registros": "Entrada 08:20 / Salida 17:00",
                        "observaciones": ["Tardanza severa (+20m sobre 08:00)"],
                        "estado": "severa",
                    },
                    {
                        "fecha": date(2026, 7, 7),
                        "llegada": "08:04",
                        "salida": "17:00",
                        "hora_programada": "08:00",
                        "almuerzo_duracion": "75m",
                        "almuerzo_salida": "12:00",
                        "almuerzo_regreso": "13:15",
                        "almuerzo_exceso": 15,
                        "tiempo_dentro": "8h 56m",
                        "n_registros": 4,
                        "detalle_registros": "Entrada 08:04 / Salida 17:00",
                        "observaciones": ["Tardanza leve (+4m sobre 08:00)", "Exceso almuerzo: +15m"],
                        "estado": "leve",
                    },
                    {
                        "fecha": date(2026, 7, 8),
                        "llegada": None,
                        "salida": None,
                        "hora_programada": "08:00",
                        "almuerzo_duracion": None,
                        "almuerzo_salida": None,
                        "almuerzo_regreso": None,
                        "almuerzo_exceso": None,
                        "tiempo_dentro": None,
                        "n_registros": 0,
                        "detalle_registros": "",
                        "observaciones": ["Ausencia no registrada"],
                        "estado": "ausente",
                    },
                    {
                        "fecha": date(2026, 7, 9),
                        "llegada": "08:10",
                        "salida": "17:00",
                        "hora_programada": "08:00",
                        "almuerzo_duracion": "60m",
                        "almuerzo_salida": "12:00",
                        "almuerzo_regreso": "13:00",
                        "almuerzo_exceso": None,
                        "tiempo_dentro": "8h 50m",
                        "n_registros": 4,
                        "detalle_registros": "",
                        "observaciones": ["Tardanza JUSTIFICADA: Cita médica"],
                        "estado": "leve",
                        "justificado": True,
                    },
                ],
                "resumen": {
                    "total_dias": 4,
                    "tardanza_severa": 1,
                    "tardanza_leve": 1,
                    "almuerzo_largo": 1,
                    "incompletos": 0,
                    "ausencias": 1,
                    "justificadas": 1,
                    "salida_anticipada_leve": 0,
                    "salida_anticipada_severa": 0,
                    "horas_trabajadas_netas": 32.5,
                    "horas_esperadas": 32.0,
                    "diferencia_horas": 0.5,
                },
                "sin_novedades": False,
                "horario_info": {
                    "tipo": "Lunes a Viernes 08:00 a 17:00",
                    "descripcion": "Turno administrativo regular",
                },
            }
        }

        filtros = dict(DEFAULT_FILTROS)
        filtros.update({
            "columna_tiempo_dentro": True,
            "verificar_horas": True,
            "mostrar_tiempo_extra": True,
            "mostrar_todos_los_dias": True,
        })

        generar_pdf_persona(
            ruta_salida=out_pdf,
            analisis_persona=analisis_persona,
            config=DEFAULT_CONFIG,
            nombre_archivo_origen="test.csv",
            filtros=filtros,
            sin_horario=["Sin Horario 1"],
        )

        assert os.path.exists(out_pdf)
        assert os.path.getsize(out_pdf) > 0


# ==============================================================================
# 4. RENDERING DOCX (python-docx)
# ==============================================================================
class TestReportDocxCoverage:

    def test_helpers_docx(self):
        doc = Document()
        _setup_document(doc)
        _add_heading(doc, "Título Test", level=1)
        _add_rule(doc)

        table = doc.add_table(rows=2, cols=2)
        cell = table.cell(0, 0)
        _set_cell_bg(cell, RGBColor(240, 240, 240))
        _set_cell_borders(cell)
        _cell_text(cell, "Texto celda", bold=True, color=RGBColor(50, 50, 50), size_pt=10)

        assert _rgb_hex(RGBColor(255, 0, 128)) == "FF0080"
        assert _fmt_fecha("2026-07-06") == "2026-07-06"
        assert _fmt_fecha(date(2026, 7, 6)) == "06/07/2026"

        _tabla_portada(doc, [["Campo", "Valor"]])
        _tabla_datos(doc, ["Col1", "Col2"], [["Val1", "Val2"]])
        _seccion_lista_tardanza(
            doc,
            "Tardanzas",
            "Sub",
            [{"nombre": "Ana", "hora": "08:10", "programado": "08:00", "retraso": 10}],
            RGBColor(255, 230, 230),
        )
        _seccion_almuerzo(doc, [{"nombre": "Ana", "salida": "12:00", "regreso": "13:15", "duracion": 75, "exceso": 15}])
        _seccion_incompletos(doc, [{"nombre": "Ana", "registros": 1, "detalle": "Solo entrada"}])
        _seccion_salidas_anticipadas(doc, "Salidas", [{"nombre": "Ana", "hora": "16:00", "programado": "17:00", "retraso": 60}], RGBColor(255, 200, 200))
        _seccion_detalle_persona(
            doc,
            [{"fecha": date(2026, 7, 6), "llegada": "08:00", "salida": "17:00", "tiempo_dentro": "9h"}],
            col_tiempo_dentro=True,
            color_novedad=RGBColor(255, 230, 230),
        )
        _seccion_horas_contrato_docx(doc, {"horas_trabajadas_netas": 40, "horas_esperadas": 40, "diferencia_horas": 0})

    def test_generar_docx_general(self, tmp_path):
        out_docx = str(tmp_path / "general.docx")
        analisis_por_dia = {
            date(2026, 7, 6): {
                "tardanza_severa": [{"nombre": "Carlos", "hora": "08:15", "programado": "08:00", "retraso": 15}],
                "tardanza_leve": [{"nombre": "Maria", "hora": "09:03", "programado": "09:00", "retraso": 3}],
                "almuerzo_largo": [{"nombre": "Carlos", "salida": "12:00", "regreso": "13:15", "duracion": 75, "exceso": 15}],
                "salida_anticipada_severa": [{"nombre": "Maria", "hora": "16:00", "programado": "18:00", "retraso": 120}],
                "salida_anticipada_leve": [],
                "registros_incompletos": [{"nombre": "Pedro", "registros": 1, "detalle": "Entrada"}],
                "resumen": {
                    "total_personas": 3,
                    "tardanza_severa": 1,
                    "tardanza_leve": 1,
                    "almuerzo_largo": 1,
                    "incompletos": 1,
                    "salida_anticipada_severa": 1,
                    "salida_anticipada_leve": 0,
                },
            }
        }
        log_dup = [{"nombre": "Carlos", "fecha": date(2026, 7, 6), "hora_orig": "08:00", "hora_dup": "08:00", "tipo": "Entrada", "diferencia": 0.1}]

        generar_docx(
            ruta_salida=out_docx,
            analisis_por_dia=analisis_por_dia,
            log_duplicados=log_dup,
            config=DEFAULT_CONFIG,
            nombre_archivo_origen="test.csv",
            filtros=DEFAULT_FILTROS,
            sin_horario=["Pedro"],
        )

        assert os.path.exists(out_docx)
        doc = Document(out_docx)
        assert len(doc.tables) > 0

    def test_generar_docx_persona(self, tmp_path):
        out_docx = str(tmp_path / "persona.docx")
        analisis_persona = {
            "Carlos Perez": {
                "dias": [
                    {
                        "fecha": date(2026, 7, 6),
                        "llegada": "08:20",
                        "salida": "17:00",
                        "hora_programada": "08:00",
                        "almuerzo_duracion": "60m",
                        "almuerzo_salida": "12:00",
                        "almuerzo_regreso": "13:00",
                        "almuerzo_exceso": None,
                        "tiempo_dentro": "8h 40m",
                        "n_registros": 4,
                        "detalle_registros": "Entrada 08:20 / Salida 17:00",
                        "observaciones": ["Tardanza severa (+20m sobre 08:00)"],
                        "estado": "severa",
                    },
                    {
                        "fecha": date(2026, 7, 7),
                        "llegada": None,
                        "salida": None,
                        "hora_programada": "08:00",
                        "almuerzo_duracion": None,
                        "almuerzo_salida": None,
                        "almuerzo_regreso": None,
                        "almuerzo_exceso": None,
                        "tiempo_dentro": None,
                        "n_registros": 0,
                        "detalle_registros": "",
                        "observaciones": ["Ausencia"],
                        "estado": "ausente",
                    },
                    {
                        "fecha": date(2026, 7, 8),
                        "llegada": "08:00",
                        "salida": "17:00",
                        "hora_programada": "08:00",
                        "almuerzo_duracion": "60m",
                        "almuerzo_salida": "12:00",
                        "almuerzo_regreso": "13:00",
                        "almuerzo_exceso": None,
                        "tiempo_dentro": "9h",
                        "n_registros": 4,
                        "detalle_registros": "",
                        "observaciones": [],
                        "estado": "ok",
                    },
                ],
                "resumen": {
                    "total_dias": 3,
                    "tardanza_severa": 1,
                    "tardanza_leve": 0,
                    "almuerzo_largo": 0,
                    "incompletos": 0,
                    "ausencias": 1,
                    "justificadas": 0,
                    "salida_anticipada_leve": 0,
                    "salida_anticipada_severa": 0,
                    "horas_trabajadas_netas": 17.6,
                    "horas_esperadas": 24.0,
                    "diferencia_horas": -6.4,
                },
                "sin_novedades": False,
                "horario_info": {
                    "tipo": "Lunes a Viernes 08:00 a 17:00",
                    "descripcion": "Turno administrativo regular",
                },
            }
        }

        filtros = dict(DEFAULT_FILTROS)
        filtros.update({
            "columna_tiempo_dentro": True,
            "verificar_horas": True,
            "mostrar_tiempo_extra": True,
            "mostrar_todos_los_dias": True,
        })

        generar_docx_persona(
            ruta_salida=out_docx,
            analisis_persona=analisis_persona,
            config=DEFAULT_CONFIG,
            nombre_archivo_origen="test.csv",
            filtros=filtros,
            sin_horario=["Pedro"],
        )

        assert os.path.exists(out_docx)
        doc = Document(out_docx)
        assert len(doc.tables) > 0


# ==============================================================================
# 5. BUILD_PDF DISPATCHER
# ==============================================================================
class TestBuildPdfDispatcher:

    def test_build_pdf_sin_horarios_falla(self, tmp_path):
        pdf_path = str(tmp_path / "dummy.pdf")
        with patch("db.get_horarios", return_value={"by_id": {}, "by_nombre": {}}):
            with pytest.raises(ValueError, match="No se pueden generar reportes sin horarios cargados"):
                build_pdf(
                    registros=[],
                    config={},
                    modo="general",
                    persona="",
                    pdf_path=pdf_path,
                    nombre_origen="dummy.csv",
                )

    def test_build_pdf_general_pdf_y_docx(self, tmp_path, mock_horarios_standard):
        pdf_path = str(tmp_path / "rep.pdf")
        docx_path = str(tmp_path / "rep.docx")
        registros = [
            _make_punch("Carlos Perez", "101", "2026-07-06", "08:00", "Entrada"),
            _make_punch("Carlos Perez", "101", "2026-07-06", "17:00", "Salida"),
        ]

        with patch("db.get_horarios", return_value=mock_horarios_standard), \
             patch("db.get_justificaciones_dict", return_value={}), \
             patch("db.get_feriados_set", return_value=set()), \
             patch("db.get_breaks_categorizados_dict", return_value={}):

            build_pdf(
                registros=registros,
                config=DEFAULT_CONFIG,
                modo="general",
                persona="",
                pdf_path=pdf_path,
                nombre_origen="test.csv",
                filtros={"mostrar_todos_los_dias": True},
                formato="pdf",
            )
            assert os.path.exists(pdf_path)

            build_pdf(
                registros=registros,
                config=DEFAULT_CONFIG,
                modo="general",
                persona="",
                pdf_path=docx_path,
                nombre_origen="test.csv",
                filtros={"mostrar_todos_los_dias": True},
                formato="docx",
            )
            assert os.path.exists(docx_path)

    def test_build_pdf_persona_pdf_y_docx(self, tmp_path, mock_horarios_standard):
        pdf_path = str(tmp_path / "per.pdf")
        docx_path = str(tmp_path / "per.docx")
        registros = [
            _make_punch("Carlos Perez", "101", "2026-07-06", "08:00", "Entrada"),
            _make_punch("Carlos Perez", "101", "2026-07-06", "17:00", "Salida"),
        ]

        with patch("db.get_horarios", return_value=mock_horarios_standard), \
             patch("db.get_justificaciones_dict", return_value={}), \
             patch("db.get_feriados_set", return_value=set()), \
             patch("db.get_breaks_categorizados_dict", return_value={}):

            build_pdf(
                registros=registros,
                config=DEFAULT_CONFIG,
                modo="persona",
                persona="Carlos Perez",
                pdf_path=pdf_path,
                nombre_origen="test.csv",
                filtros={"mostrar_todos_los_dias": True},
                formato="pdf",
            )
            assert os.path.exists(pdf_path)

            build_pdf(
                registros=registros,
                config=DEFAULT_CONFIG,
                modo="persona",
                persona="Carlos Perez",
                pdf_path=docx_path,
                nombre_origen="test.csv",
                filtros={"mostrar_todos_los_dias": True},
                formato="docx",
            )
            assert os.path.exists(docx_path)


# ==============================================================================
# 6. MISC UTILS EN REPORTS
# ==============================================================================
class TestReportsMiscUtils:

    def test_fmt_horas(self):
        assert _fmt_horas(None) == "—"
        assert _fmt_horas(0) == "0m"
        assert _fmt_horas(45) == "45m"
        assert _fmt_horas(90) == "1h 30m"
        assert _fmt_horas(-90) == "-1h 30m"

    def test_retraso_de_obs(self):
        assert _retraso_de_obs(["Tardanza (+15m)"]) == "+15 min"
        assert _retraso_de_obs(["Sin novedades"]) == "—"
